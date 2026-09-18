"""
Build the delivered .xlsx workbook: raw enriched question rows plus
pivot-style summary sheets, from the CSVs build.py / build_graph.py produced.
"""
import csv
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PROC = ROOT / "data" / "processed"
OUT_XLSX = ROOT / "Lok_Rajya_Sabha_PQ_Analysis.xlsx"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1D1A14")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT_NAME, size=10)


def df(name):
    path = PROC / name
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path)


def main():
    sheets = {
        "Party Ranking": df("party_summary.csv"),
        "Ministry Ranking": df("ministry_summary.csv"),
        "State Ranking": df("state_summary.csv"),
        "Party x Ministry": df("party_ministry_matrix.csv"),
        "Party x State": df("state_party_matrix.csv"),
        "Party x Type": df("party_type_matrix.csv"),
        "Topic Keywords": df("topic_keywords.csv"),
        "KG Signature Ministries": df("kg_party_ministry_lift.csv"),
        "KG Party Communities": df("kg_party_communities.csv"),
        "KG Ministry Co-occurrence": df("kg_ministry_cooccurrence.csv"),
        "KG Cross-Party Bridges": df("kg_mp_coasking_bridges.csv"),
        "PIB Ministry Comparison": df("pib_ministry_comparison.csv"),
        "PIB Scheme Comparison": df("pib_scheme_comparison.csv"),
        "PLI Scheme Scrutiny": df("pli_scheme_scrutiny.csv"),
        "Ministry Scheme Scrutiny": df("ministry_scheme_scrutiny.csv"),
        "PARIVESH State Comparison": df("parivesh_state_comparison.csv"),
        "PARIVESH Activity Summary": df("parivesh_activity_summary.csv"),
        "PARIVESH Mega Projects": df("parivesh_mega_projects.csv"),
        "Unmatched Members": df("unmatched_members.csv"),
        "LS Questions (raw)": df("ls_questions_enriched.csv"),
        "RS Questions (raw)": df("rs_questions_enriched.csv"),
    }

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name[:31], index=False)

    wb = load_workbook(OUT_XLSX)

    # Cover / README sheet, inserted first
    cover = wb.create_sheet("Read Me", 0)
    cover_lines = [
        ("Lok & Rajya Sabha — Parliament Questions Analysis", True, 14),
        ("", False, 10),
        ("Coverage: 18th Lok Sabha since formation (June 2024); Rajya Sabha sessions 265-271", False, 10),
        ("(the sessions spanning the same period).", False, 10),
        ("", False, 10),
        ("Sources: sansad.in (api_ls question & member APIs), rsdoc.nic.in (Rajya Sabha", False, 10),
        ("question search). Party attribution cross-references both houses' member rosters;", False, 10),
        ("names reconciled across the chambers' independently-maintained abbreviation", False, 10),
        ("conventions (e.g. Shiv Sena (UBT), YSRCP, NCP (SP)).", False, 10),
        ("", False, 10),
        ("Sheet guide:", True, 11),
        ("  Party / Ministry / State Ranking — totals, sorted descending.", False, 10),
        ("  Party x Ministry / State / Type — cross-tabs (pivot-ready).", False, 10),
        ("  Topic Keywords — subject-line keyword frequency, overall (party=ALL) and per party.", False, 10),
        ("  KG * sheets — knowledge-graph analysis: signature ministries (lift = how much more", False, 10),
        ("    a party asks about a ministry than its overall question share predicts), Louvain", False, 10),
        ("  community clusters, ministry co-occurrence, and cross-party MP co-asking pairs.", False, 10),
        ("  Unmatched Members — the 138 LS questions (0.4%) whose asking MP's name could not", False, 10),
        ("    be matched to the member roster; excluded from party/state breakdowns.", False, 10),
        ("  PIB Ministry/Scheme Comparison — PQ volume vs PIB (Press Information Bureau) press", False, 10),
        ("    releases over the same window: which ministries/schemes get heavy parliamentary", False, 10),
        ("    scrutiny relative to how much they self-publicize, and vice versa. See caveats in", False, 10),
        ("    docs/ANALYSIS.md — the External Affairs row is a known PIB-index gap, not reality,", False, 10),
        ("    and scheme matching is title-text-only (a floor on PIB coverage, not a ceiling).", False, 10),
        ("  PLI Scheme Scrutiny — PLI sub-scheme A-F disbursal grades (from a sibling repo's", False, 10),
        ("    PIB-sourced research, not re-verified here) joined against how often each scheme", False, 10),
        ("    is specifically named in PQ text: worst-performing schemes vs how much scrutiny", False, 10),
        ("    they actually draw.", False, 10),
        ("  Ministry Scheme Scrutiny — MeitY/DoT/DPIIT's full scheme catalogues (live WordPress", False, 10),
        ("    APIs behind their JS-shell sites), matched against PQ text, with months since each", False, 10),
        ("    ministry's own CMS last touched a scheme's page. Keyword-heuristic match; lower bound.", False, 10),
        ("  PARIVESH sheets — Environmental Clearance proposals (live from parivesh.nic.in) vs.", False, 10),
        ("    Environment-ministry PQs, by state; activity-category volume; and the 25 largest", False, 10),
        ("    proposals by investment, checked for whether the applicant company is named in a PQ.", False, 10),
        ("  LS/RS Questions (raw) — one row per question, enriched with party/state.", False, 10),
        ("    question_text/answer_text are populated for Rajya Sabha only — the Lok Sabha API", False, 10),
        ("    used here doesn't expose full text, only subject lines (see repo README).", False, 10),
        ("", False, 10),
        ("Full methodology, the graph-analysis scripts, the DuckDB database and the interactive", False, 10),
        ("dashboard are in the repository this workbook ships with.", False, 10),
    ]
    for i, (text, bold, size) in enumerate(cover_lines, start=1):
        c = cover.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT_NAME, bold=bold, size=size)
    cover.column_dimensions["A"].width = 100

    for ws in wb.worksheets:
        if ws.title == "Read Me":
            continue
        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = BODY_FONT
        for col_cells in ws.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells[:200])
            col_letter = get_column_letter(col_cells[0].column)
            ws.column_dimensions[col_letter].width = min(max(length + 2, 9), 42)
        ws.auto_filter.ref = ws.dimensions

    wb.save(OUT_XLSX)
    print("Wrote", OUT_XLSX, OUT_XLSX.stat().st_size, "bytes")


if __name__ == "__main__":
    main()
